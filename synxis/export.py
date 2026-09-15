import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from . import config
from .client import build_session, fetch_range, date_for_calendar_id


def run():
    session = build_session()
    print(f"Exporting current rates {config.EXPORT_START_DATE} -> {config.EXPORT_END_DATE} ...")
    rows, room_map, _ = fetch_range(session, config.EXPORT_START_DATE, config.EXPORT_END_DATE)

    if not rows:
        print("No rates returned. Check that SYNXIS_COOKIE / SYNXIS_CSRF_TOKEN are fresh.")
        return

    wanted = {c.strip().upper() for c in config.EXPORT_RATE_CODES}
    records = [
        {
            "Date": date_for_calendar_id(r["calendar_id"]).isoformat(),
            "Room Type": r["room_name"],
            "Rate Code": r["rate_code"],
            "Current Rate": r["value"],
            "New Rate": None,
        }
        for r in sorted(rows, key=lambda x: (x["rate_code"], x["room_name"], x["calendar_id"]))
        if not wanted or r["rate_code"].strip().upper() in wanted
    ]

    df = pd.DataFrame(records, columns=["Date", "Room Type", "Rate Code", "Current Rate", "New Rate"])
    try:
        df.to_excel(config.EXCEL_PATH, sheet_name=config.SHEET_NAME, index=False)
    except PermissionError:
        print(f"Cannot write {config.EXCEL_PATH} — close it in Excel and try again.")
        return

    _format_workbook(config.EXCEL_PATH, config.SHEET_NAME)

    n_days = (config.EXPORT_END_DATE - config.EXPORT_START_DATE).days + 1
    print(f"Wrote {len(df)} rows ({len(room_map)} room types, "
          f"{len(df['Rate Code'].unique())} rate code(s), {n_days} days) to {config.EXCEL_PATH}")
    print("Fill the yellow 'New Rate' column, save, then run: python main.py push")


def _format_workbook(path: str, sheet_name: str):
    """Highlights the editable 'New Rate' column and styles the header row."""
    wb = load_workbook(path)
    ws = wb[sheet_name]
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4472C4")
    yellow = PatternFill("solid", fgColor="FFFF00")

    for col in range(1, 6):
        c = ws.cell(row=1, column=col)
        c.font, c.fill, c.alignment = header_font, header_fill, Alignment(horizontal="center")

    for col_letter, width in {"A": 12, "B": 34, "C": 14, "D": 13, "E": 13}.items():
        ws.column_dimensions[col_letter].width = width

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=5).fill = yellow  # New Rate column

    ws.cell(row=1, column=7, value="Fill only the yellow 'New Rate' column. "
                                    "Leave a row blank to keep its current rate.")
    wb.save(path)
